# -*- coding: utf-8 -*-
"""
把《医保待测交接表》里用户更新的状态，同步成总表可用的「人工覆盖」文件。

为什么要这么绕：总表副本 `..._本轮已验_<日期>.xlsx` 是 `update_tracking_excel.py`
**每次从源文件重新生成**的，直接改副本会被下一轮覆盖。
所以同步走 `reports/verify/manual_status.json`，由 `update_tracking_excel.py` 在写回时优先套用。
"""
import json
from pathlib import Path

import openpyxl

INS = Path(r"D:\workapp\ZJCK\HIS_EMR上线测试用例追踪表_医保待测_交接_20261009.xlsx")
OUT = Path(__file__).resolve().parent / "reports" / "verify" / "manual_status.json"

# 只认这些"原表的词"
VALID = {"通过", "失败", "阻塞", "跳过", "未执行"}


def main() -> int:
    wb = openpyxl.load_workbook(INS)
    ws = wb["用例清单"]
    overrides = {}
    for r in range(2, ws.max_row + 1):
        cid = str(ws.cell(row=r, column=2).value or "").strip()
        st = str(ws.cell(row=r, column=13).value or "").strip()
        if not cid or st not in VALID:
            continue
        if st == "未执行":
            note = "医保功能已对接，待验证"
        else:
            note = f"随《医保待测交接表》同步（用户更新）：{st}"
        overrides[cid] = {"status": st, "note": note}
    wb.close()
    OUT.write_text(json.dumps(overrides, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"[产出] {OUT}")
    print(f"  共 {len(overrides)} 条：")
    for k, v in overrides.items():
        print(f"    {k:8s} {v['status']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
