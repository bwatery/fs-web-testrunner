# -*- coding: utf-8 -*-
"""
导出「依赖打印设备/窗口外设、本环境验不了」的用例为独立 Excel（保持源表格式），
交给有打印机的同事去测。

做法与 `export_insurance_excel.py` 一致：复制源文件 -> 在"用例清单"里删掉不需要的行
-> 备注写清依赖 -> 序号重编 -> 去掉与本次交接无关的工作表 -> 另存。

**与医保交接表的重叠**：L6-063 / L6-068 / L7-004 既需医保前置机又需打印设备，
已在《医保待测_交接》里交出去了，本表**不再重复**，只在控制台提示。
"""
import shutil
from pathlib import Path

import openpyxl

SRC = Path(r"D:\workapp\ZJCK\HIS_EMR上线测试用例追踪表_合并版.xlsx")
DST = Path(r"D:\workapp\ZJCK\HIS_EMR上线测试用例追踪表_打印待测_交接_20261009.xlsx")

PRINT_HW = "需打印设备与打印控件（针式/热敏/激光）"
PERIPH = "需窗口外设硬件（读卡器/扫码枪/高拍仪/叫号屏）"

# 保留的用例 → 为什么本环境验不了
KEEP = {
    "L0-013": f"{PRINT_HW}；本条本身就是「打印模板配置与单据打印」",
    "L2-030": "标本条码需条码打印机与绑定工位",
    "L2-091": f"{PRINT_HW}（申请单预览与打印）",
    "L2-100": f"{PRINT_HW}（病案首页打印导出）",
    "L2-103": f"{PRINT_HW}（传染病报告卡打印）",
    "L3-004": "需条码打印机（护士站侧标本条码打印与绑定）",
    "L3-010": f"{PRINT_HW}（检验报告打印）",
    "L3-013": f"{PRINT_HW}（检查申请单打印）",
    "L4-017": "需腕带打印机与扫码枪（执行防错扫码）",
    "L4-057": f"{PRINT_HW}（手术申请单打印）",
    "L4-060": f"{PRINT_HW}（医嘱单打印与状态回写）",
    "L5-029": f"{PRINT_HW}（病历打印与复印管理）",
    "L6-005": f"本环境已验证挂号列表联动部分并通过；仅「补打条码」{PRINT_HW}",
    "L6-013": f"{PRINT_HW}（挂号票打印与补打）",
    "L6-035": f"票据号段管理本环境可验；「发票重打」{PRINT_HW}",
    "L6-040": f"{PRINT_HW}（患者日清单生成与打印）",
    "L6-050": f"本环境已验证预交金日结部分并通过；仅「收据补打」{PRINT_HW}",
    "L6-053": f"取消结算部分本环境可验；「结算单打印」{PRINT_HW}",
    "L6-064": f"单据配置维护本环境可验；「路由调起打印」{PRINT_HW}",
    "L6-065": f"{PRINT_HW}（挂号票/指引单/收费小票）",
    "L6-066": f"{PRINT_HW}（门诊费用明细清单与取药凭证）",
    "L6-067": f"{PRINT_HW}（日清单/催款单/结算单）",
    "L8-003": f"{PRINT_HW}；本条本身就是「打印控件与打印机适配」",
}
# 已在《医保待测_交接》里交出去、本表不重复
ALREADY_HANDED = {"L6-063", "L6-068", "L7-004"}


def main() -> int:
    if not SRC.exists():
        print(f"[致命] 源文件不存在: {SRC}")
        return 2
    shutil.copyfile(SRC, DST)

    wb = openpyxl.load_workbook(DST)
    ws = wb["用例清单"]

    keep_rows, cids = [], []
    for r in range(2, ws.max_row + 1):
        cid = str(ws.cell(row=r, column=2).value or "").strip()
        if cid in KEEP:
            keep_rows.append(r)
            cids.append(cid)

    keep = set(keep_rows)
    for r in range(ws.max_row, 1, -1):
        if r not in keep:
            ws.delete_rows(r)

    for r in range(2, ws.max_row + 1):
        cid = str(ws.cell(row=r, column=2).value or "").strip()
        old = str(ws.cell(row=r, column=17).value or "").strip()
        note = f"本环境未验：{KEEP[cid]}"
        if old and old not in note:
            note = f"{note}（原备注：{old}）"
        ws.cell(row=r, column=17, value=note)
        ws.cell(row=r, column=1, value=r - 1)          # 序号重编 1..N

    # 门禁看板是全项目公式看板，在只剩 23 行的副本里会算出误导数字；
    # 缺陷追踪是全项目缺陷清单，与本次交接无关 —— 一并去掉。
    for name in ("门禁看板", "缺陷追踪"):
        if name in wb.sheetnames:
            del wb[name]

    wb.save(DST)
    wb.close()

    print(f"导出 {len(cids)} 条：{', '.join(cids)}")
    print(f"文件：{DST}")
    print(f"（已跳过、不重复交接：{', '.join(sorted(ALREADY_HANDED))} —— 在《医保待测_交接》里）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
